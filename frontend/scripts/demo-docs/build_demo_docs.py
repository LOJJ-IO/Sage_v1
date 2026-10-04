"""Builds the `/demo` store documents (real PDFs + Word files) from one source.

Outputs, all generated — edit this file, not the outputs:
  public/demo/*.pdf, public/demo/*.docx      what the preview viewers render
  src/lib/demo/documents.generated.ts         plain text of each doc, used by the
                                              mock /ask for citation offsets

Every quote in `DEMO_ANSWERS` (src/lib/demo/seed.ts) must appear verbatim in a
doc's text — `mock-api.test.ts` enforces that against the generated file.

Run from `frontend/`:
  pip install python-docx
  python scripts/demo-docs/build_demo_docs.py
PDFs are printed with headless Chrome; set CHROME=/path/to/chrome if it isn't
at the default macOS location.
"""

from __future__ import annotations

import html
import json
import os
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
PUBLIC_DIR = ROOT / "public" / "demo"
GENERATED_TS = ROOT / "src" / "lib" / "demo" / "documents.generated.ts"
CHROME = os.environ.get(
    "CHROME",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
)

STORE = "Juniper Lane Boutique"
ACCENT = "3F6B4F"  # sage green
ACCENT_SOFT = "EEF4EF"
INK = "1F2A24"
MUTED = "6B7570"

# ── content ────────────────────────────────────────────────────────────────
# Block kinds: h2, p, ul, ol, checklist, table, callout.


@dataclass
class Doc:
    file_id: str
    filename: str
    title: str
    subtitle: str
    meta: list[tuple[str, str]]
    blocks: list[tuple]


DOCS: list[Doc] = [
    Doc(
        file_id="demo-open-close",
        filename="Opening-and-Closing-Checklist.pdf",
        title="Opening & Closing Checklist",
        subtitle="Daily routine for the shift lead on duty",
        meta=[("Effective", "March 1, 2026"), ("Owner", "Maya Thompson, Store Manager"), ("Version", "3.2")],
        blocks=[
            ("p", "Use this checklist every day. Initial each section in the shift log once it's done, and tell a manager straight away if anything can't be completed."),
            ("h2", "Opening — arrive 30 minutes before doors"),
            ("checklist", [
                "Disarm the alarm — the code is on the key fob card in the manager's lockbox.",
                "Lights on: front panel first, then fitting rooms.",
                "Count the float: $200 in the till — $100 in fives and tens, $100 in coins and ones.",
                "Steam any new arrivals on the \"to floor\" rack.",
                "Check the shared inbox for online pickup orders and pull them to the hold shelf.",
            ]),
            ("table", ["Till float", "Amount"], [
                ["Fives and tens", "$100"],
                ["Coins and ones", "$100"],
                ["Total float", "$200"],
            ]),
            ("h2", "Closing — start 15 minutes before doors close"),
            ("checklist", [
                "Last call in fitting rooms at 10 minutes to close.",
                "Run the end-of-day report on the POS and print two copies.",
                "Count the till down to the $200 float; the rest goes in the deposit bag.",
                "Deposit bag goes in the back safe — never leave cash in the till overnight.",
                "Return fitting room items to the floor, lights off, set the alarm.",
            ]),
            ("callout", "If the till is over or short", "Recount once with a second staff member. If it's still off by more than $5, note the amount on both end-of-day copies and text the store manager before you leave."),
        ],
    ),
    Doc(
        file_id="demo-pos-troubleshooting",
        filename="POS-Troubleshooting-Guide.pdf",
        title="POS Troubleshooting Guide",
        subtitle="Quick fixes for the register, card reader and printer",
        meta=[("Effective", "January 15, 2026"), ("Owner", "Jordan Lee, Assistant Manager"), ("Version", "1.4")],
        blocks=[
            ("p", "Try the fixes below in order. Never turn a customer away because of a technical problem — there is always a backup way to take payment."),
            ("h2", "Card reader not connecting"),
            ("ol", [
                "Unplug the reader's USB cable, wait 10 seconds, plug it back in.",
                "If it still says \"offline\", restart the iPad.",
                "Take cash or e-transfer and call the POS support line on the sticker under the counter.",
            ]),
            ("h2", "Receipt printer jammed"),
            ("ol", [
                "Open the lid and pull the paper roll out.",
                "Tear off the crumpled section.",
                "Reload the roll with the paper feeding from the bottom.",
            ]),
            ("h2", "Price won't scan"),
            ("p", "Key in the SKU from the tag manually. If the item isn't in the system, ask a manager before selling it."),
            ("table", ["Problem", "Backup payment", "Who to call"], [
                ["Card reader offline", "Cash or e-transfer", "POS support line"],
                ["iPad won't start", "Cash only, handwritten receipt", "Store manager"],
                ["Internet down", "Card reader offline mode (up to $500)", "Store manager"],
            ]),
            ("callout", "Offline card payments", "In offline mode the reader stores up to $500 of card payments and sends them once the internet is back. Keep the iPad plugged in until they've gone through."),
        ],
    ),
    Doc(
        file_id="demo-employee-discount",
        filename="Employee-Discount-Policy.docx",
        title="Employee Discount Policy",
        subtitle="Who gets a discount, how much, and how to ring it up",
        meta=[("Effective", "February 1, 2026"), ("Owner", "Maya Thompson, Store Manager"), ("Version", "2.0")],
        blocks=[
            ("h2", "Discount rates"),
            ("p", "All staff get 30% off full-price items and 10% off sale items after their first 30 days."),
            ("table", ["Item type", "Discount", "Starts"], [
                ["Full-price items", "30%", "After 30 days"],
                ["Sale items", "10%", "After 30 days"],
                ["Gift cards", "None", "—"],
            ]),
            ("h2", "Who can use it"),
            ("p", "The discount is for the employee and one immediate family member only. Ring it up under your own employee number — never another staff member's."),
            ("h2", "Ringing it up"),
            ("ul", [
                "Employee purchases must be rung up by a different staff member, not yourself.",
                "The discount cannot be combined with other promotions or used on gift cards.",
                "Keep the receipt with the item if you're taking it home at the end of a shift.",
            ]),
            ("callout", "Misuse", "Using the discount for friends or ringing up your own purchase can lead to losing the discount or other disciplinary action."),
        ],
    ),
    Doc(
        file_id="demo-onboarding",
        filename="New-Hire-Onboarding-Guide.docx",
        title="New Hire Onboarding Guide",
        subtitle="Your first week at Juniper Lane",
        meta=[("Effective", "April 1, 2026"), ("Owner", "Jordan Lee, Assistant Manager"), ("Version", "1.1")],
        blocks=[
            ("p", "Welcome to the team! This guide walks you through your first week so you know what to expect and who to ask."),
            ("h2", "Your first day"),
            ("p", "On your first day, arrive at 9:30am and ask for the shift lead on duty. Bring two pieces of ID and your banking details for payroll."),
            ("h2", "First-week schedule"),
            ("table", ["Day", "Focus", "With"], [
                ["Day 1", "Store tour, policies, POS login", "Shift lead"],
                ["Day 2", "Fitting rooms and floor recovery", "Senior associate"],
                ["Day 3", "Receiving new stock and steaming", "Senior associate"],
                ["Day 4", "Till practice and returns", "Shift lead"],
                ["Day 5", "First solo floor shift", "Check-in with manager"],
            ]),
            ("h2", "Training"),
            ("ul", [
                "You'll shadow a senior associate for your first three shifts before working the till on your own.",
                "Read the Returns & Exchanges Policy and the Opening & Closing Checklist before Day 4.",
                "Ask Sage any policy question — it answers from our own store documents.",
            ]),
            ("callout", "Dress code", "Wear solid neutral colours (black, white, cream or grey) and closed-toe shoes. Name tags are kept in the back office."),
        ],
    ),
]

# ── plain text (mock /ask + citations) ─────────────────────────────────────


def to_plain_text(doc: Doc) -> str:
    lines = [doc.title, doc.subtitle, " · ".join(f"{k}: {v}" for k, v in doc.meta), ""]
    for block in doc.blocks:
        kind = block[0]
        if kind == "h2":
            lines += [block[1], ""]
        elif kind == "p":
            lines += [block[1], ""]
        elif kind == "ul":
            lines += [f"- {item}" for item in block[1]] + [""]
        elif kind == "ol":
            lines += [f"{i}. {item}" for i, item in enumerate(block[1], 1)] + [""]
        elif kind == "checklist":
            lines += [f"[ ] {item}" for item in block[1]] + [""]
        elif kind == "table":
            lines += [" | ".join(block[1])] + [" | ".join(row) for row in block[2]] + [""]
        elif kind == "callout":
            lines += [f"{block[1]}: {block[2]}", ""]
    return "\n".join(lines).rstrip() + "\n"


# ── PDF (HTML → headless Chrome) ───────────────────────────────────────────

CSS = f"""
@page {{ size: Letter; margin: 0.75in 0.8in 0.9in; }}
* {{ box-sizing: border-box; }}
body {{ font-family: -apple-system, "Helvetica Neue", Arial, sans-serif; color: #{INK};
  font-size: 10.5pt; line-height: 1.5; margin: 0; }}
.letterhead {{ display: flex; align-items: center; gap: 10px; padding-bottom: 14px;
  border-bottom: 2px solid #{ACCENT}; margin-bottom: 22px; }}
.mark {{ width: 34px; height: 34px; border-radius: 50%; background: #{ACCENT}; color: white;
  display: flex; align-items: center; justify-content: center; font: 600 12pt Georgia, serif; }}
.store {{ font-size: 8.5pt; letter-spacing: .14em; text-transform: uppercase; color: #{MUTED}; }}
.label {{ margin-left: auto; font-size: 8pt; letter-spacing: .1em; text-transform: uppercase;
  color: #{ACCENT}; border: 1px solid #{ACCENT}; border-radius: 999px; padding: 2px 10px; }}
h1 {{ font: 600 25pt/1.15 Georgia, "Times New Roman", serif; margin: 0 0 4px; }}
.subtitle {{ color: #{MUTED}; font-size: 11.5pt; margin: 0 0 14px; }}
.meta {{ display: flex; gap: 26px; font-size: 8.5pt; color: #{MUTED}; margin-bottom: 22px; }}
.meta b {{ display: block; color: #{INK}; font-weight: 600; font-size: 9.5pt; }}
h2 {{ font: 600 13.5pt Georgia, serif; color: #{ACCENT}; margin: 22px 0 8px;
  padding-bottom: 4px; border-bottom: 1px solid #DCE3DE; break-after: avoid; }}
p {{ margin: 0 0 10px; }}
ul, ol {{ margin: 0 0 12px; padding-left: 20px; }} li {{ margin: 3px 0; }}
.checklist {{ list-style: none; padding: 0; }}
.checklist li {{ display: flex; gap: 10px; padding: 7px 0; border-bottom: 1px dashed #DCE3DE; }}
.box {{ flex: none; width: 12px; height: 12px; margin-top: 4px; border: 1.5px solid #{ACCENT}; border-radius: 3px; }}
table {{ width: 100%; border-collapse: collapse; margin: 12px 0 16px; font-size: 9.5pt; break-inside: avoid; }}
th {{ text-align: left; background: #{ACCENT}; color: white; font-weight: 600; padding: 7px 10px; }}
td {{ padding: 7px 10px; border-bottom: 1px solid #DCE3DE; }}
tr:nth-child(even) td {{ background: #F6F8F6; }}
.callout {{ background: #{ACCENT_SOFT}; border-left: 4px solid #{ACCENT}; border-radius: 4px;
  padding: 10px 14px; margin: 16px 0; break-inside: avoid; }}
.callout b {{ display: block; color: #{ACCENT}; margin-bottom: 2px; }}
footer {{ position: fixed; bottom: -0.55in; left: 0; right: 0; display: flex; justify-content: space-between;
  font-size: 7.5pt; color: #{MUTED}; border-top: 1px solid #DCE3DE; padding-top: 6px; }}
"""


def block_html(block: tuple) -> str:
    e = html.escape
    kind = block[0]
    if kind == "h2":
        return f"<h2>{e(block[1])}</h2>"
    if kind == "p":
        return f"<p>{e(block[1])}</p>"
    if kind in ("ul", "ol"):
        return f"<{kind}>" + "".join(f"<li>{e(i)}</li>" for i in block[1]) + f"</{kind}>"
    if kind == "checklist":
        return '<ul class="checklist">' + "".join(
            f'<li><span class="box"></span><span>{e(i)}</span></li>' for i in block[1]
        ) + "</ul>"
    if kind == "table":
        head = "".join(f"<th>{e(h)}</th>" for h in block[1])
        rows = "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in block[2])
        return f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>"
    if kind == "callout":
        return f'<div class="callout"><b>{e(block[1])}</b>{e(block[2])}</div>'
    raise ValueError(kind)


def doc_html(doc: Doc) -> str:
    e = html.escape
    meta = "".join(f"<div>{e(k)}<b>{e(v)}</b></div>" for k, v in doc.meta)
    body = "".join(block_html(b) for b in doc.blocks)
    version = dict(doc.meta).get("Version", "")
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{e(doc.title)}</title>
<style>{CSS}</style></head><body>
<div class="letterhead"><div class="mark">JL</div><div class="store">{e(STORE)}</div>
<div class="label">Internal · Staff only</div></div>
<h1>{e(doc.title)}</h1><p class="subtitle">{e(doc.subtitle)}</p><div class="meta">{meta}</div>
{body}
<footer><span>{e(STORE)} · {e(doc.title)}</span><span>Version {e(version)}</span></footer>
</body></html>"""


def build_pdf(doc: Doc, out: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.html"
        src.write_text(doc_html(doc), encoding="utf-8")
        subprocess.run(
            [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
             f"--user-data-dir={tmp}/profile", f"--print-to-pdf={out}", src.as_uri()],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )


# ── DOCX (python-docx) ─────────────────────────────────────────────────────


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), hex_color)
    props.append(fill)


def bottom_border(paragraph, hex_color: str, size: int = 12) -> None:
    props = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", str(size)), ("space", "4"), ("color", hex_color)):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    props.append(borders)


def set_cell_padding(table, twips: int = 100) -> None:
    props = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(twips if side in ("left", "right") else twips // 2))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    props.append(margins)


def drop_theme_color(style) -> None:
    """Word's built-in heading styles carry a theme colour that renderers
    (docx-preview, Pages) prefer over the explicit RGB — strip it."""
    rpr = style.element.rPr
    if rpr is None:
        return
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            color.attrib.pop(qn(attr), None)


def pin_fonts(rpr, font: str) -> None:
    """Explicit font on every script slot, no theme fonts — browsers render
    theme fonts (Calibri/Cambria) as Times when they aren't installed."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)


def add_list_item(d, marker: str, text: str) -> None:
    """Plain-text bullet/number with a hanging indent — Word's List Bullet uses a
    Symbol-font glyph that browsers show as an empty box."""
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-12)
    p.paragraph_format.space_after = Pt(3)
    m = p.add_run(f"{marker}\t")
    m.font.color.rgb = rgb(ACCENT)
    m.bold = True
    p.add_run(text)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(18))


def build_docx(doc: Doc, out: Path) -> None:
    d = Document()
    defaults = d.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    pin_fonts(defaults, "Arial")
    normal = d.styles["Normal"]
    # Fonts every OS ships, so the preview matches Word (Calibri falls back to Times on macOS).
    normal.font.name = "Arial"
    pin_fonts(normal.element.get_or_add_rPr(), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 26), ("Heading 1", 15)):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = rgb(INK if name == "Title" else ACCENT)
        drop_theme_color(style)
        pin_fonts(style.element.get_or_add_rPr(), "Georgia")
        # The linked character style ("Heading 1 Char") also carries theme colour/fonts.
        linked = d.styles[f"{name} Char"]
        linked.font.color.rgb = style.font.color.rgb
        drop_theme_color(linked)
        pin_fonts(linked.element.get_or_add_rPr(), "Georgia")
    d.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    d.styles["Heading 1"].paragraph_format.space_after = Pt(6)

    section = d.sections[0]
    header = section.header.paragraphs[0]
    header.text = f"{STORE.upper()}   ·   INTERNAL — STAFF ONLY"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = rgb(MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = f"{doc.title} · Version {dict(doc.meta).get('Version', '')}"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = rgb(MUTED)

    title = d.add_paragraph(doc.title, style="Title")
    bottom_border(title, ACCENT, 16)
    subtitle = d.add_paragraph()
    run = subtitle.add_run(doc.subtitle)
    run.font.size = Pt(12)
    run.font.color.rgb = rgb(MUTED)
    meta = d.add_paragraph()
    for i, (key, value) in enumerate(doc.meta):
        k = meta.add_run(("     " if i else "") + f"{key}: ")
        k.font.size = Pt(9)
        k.font.color.rgb = rgb(MUTED)
        v = meta.add_run(value)
        v.font.size = Pt(9)
        v.bold = True

    for block in doc.blocks:
        kind = block[0]
        if kind == "h2":
            d.add_heading(block[1], level=1)
        elif kind == "p":
            d.add_paragraph(block[1])
        elif kind in ("ul", "ol"):
            for i, item in enumerate(block[1], 1):
                add_list_item(d, "•" if kind == "ul" else f"{i}.", item)
        elif kind == "checklist":
            for item in block[1]:
                add_list_item(d, "☐", item)
        elif kind == "table":
            header_row, rows = block[1], block[2]
            table = d.add_table(rows=1 + len(rows), cols=len(header_row))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_cell_padding(table)
            for c, text in enumerate(header_row):
                cell = table.rows[0].cells[c]
                cell.text = text
                shade(cell, ACCENT)
                r = cell.paragraphs[0].runs[0]
                r.bold = True
                r.font.color.rgb = rgb("FFFFFF")
            for r_i, row in enumerate(rows, 1):
                for c, text in enumerate(row):
                    cell = table.rows[r_i].cells[c]
                    cell.text = text
                    if r_i % 2 == 0:
                        shade(cell, "F6F8F6")
            d.add_paragraph()
        elif kind == "callout":
            box = d.add_table(rows=1, cols=1)
            set_cell_padding(box, 160)
            cell = box.rows[0].cells[0]
            shade(cell, ACCENT_SOFT)
            p = cell.paragraphs[0]
            head = p.add_run(block[1] + "\n")
            head.bold = True
            head.font.color.rgb = rgb(ACCENT)
            p.add_run(block[2])
            d.add_paragraph()
    d.core_properties.title = doc.title
    d.core_properties.author = STORE
    d.save(out)


# ── main ───────────────────────────────────────────────────────────────────


def main() -> None:
    PUBLIC_DIR.mkdir(parents=True, exist_ok=True)
    texts: dict[str, dict[str, str]] = {}
    for doc in DOCS:
        out = PUBLIC_DIR / doc.filename
        (build_pdf if doc.filename.endswith(".pdf") else build_docx)(doc, out)
        texts[doc.file_id] = {"filename": doc.filename, "text": to_plain_text(doc)}
        print(f"built {out.relative_to(ROOT)}")

    GENERATED_TS.write_text(
        "// Generated by scripts/demo-docs/build_demo_docs.py — do not edit by hand.\n"
        "/** Plain text of each generated demo PDF/DOCX, keyed by file_id. */\n"
        f"export const GENERATED_DEMO_DOCUMENTS: Record<string, {{ filename: string; text: string }}> = {json.dumps(texts, indent=2, ensure_ascii=False)};\n",
        encoding="utf-8",
    )
    print(f"wrote {GENERATED_TS.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
